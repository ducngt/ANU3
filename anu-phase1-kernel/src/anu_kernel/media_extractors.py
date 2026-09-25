from __future__ import annotations

import io
import json
import mimetypes
import subprocess
import tempfile
import wave
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from PIL import Image
from docx import Document
from pypdf import PdfReader

from .ingestion_contracts import ExtractionStatus


@dataclass(frozen=True)
class ExtractionResult:
    media_type: str
    status: ExtractionStatus
    text: str | None
    metadata: dict[str, Any]
    analyzer_ref: str
    analyzer_version: str = "1.0.0"


def _detect(filename: str, declared: str | None, data: bytes) -> str:
    lower = filename.lower()
    # Prefer observable content signatures for formats with stable magic.
    if data.startswith(b"%PDF-"):
        return "application/pdf"
    if data.startswith(b"PK\x03\x04") and lower.endswith(".docx"):
        return "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    guessed, _ = mimetypes.guess_type(filename)
    if guessed and (guessed.startswith("image/") or guessed.startswith("audio/") or guessed.startswith("video/") or guessed.startswith("text/")):
        return guessed
    if declared and declared != "application/octet-stream":
        return declared.split(";", 1)[0].strip().lower()
    if guessed:
        return guessed
    return "application/octet-stream"


def _ffprobe(path: Path) -> dict[str, Any]:
    try:
        p = subprocess.run(
            [
                "ffprobe", "-v", "error", "-show_entries",
                "format=duration,format_name,size:stream=index,codec_type,codec_name,width,height,sample_rate,channels",
                "-of", "json", str(path),
            ],
            capture_output=True,
            text=True,
            timeout=10,
            check=False,
        )
    except (FileNotFoundError, subprocess.TimeoutExpired):
        return {}
    if p.returncode != 0:
        return {}
    try:
        return json.loads(p.stdout)
    except json.JSONDecodeError:
        return {}


def extract(filename: str, declared_media_type: str | None, data: bytes, stored_path: Path | None = None) -> ExtractionResult:
    media_type = _detect(filename, declared_media_type, data)
    base_meta: dict[str, Any] = {"filename": filename}
    try:
        if media_type == "application/pdf":
            reader = PdfReader(io.BytesIO(data))
            text = "\n".join((page.extract_text() or "") for page in reader.pages).strip()
            base_meta["page_count"] = len(reader.pages)
            return ExtractionResult(media_type, ExtractionStatus.EXTRACTED, text or None, base_meta, "urn:anu:analyzer:pdf:pypdf")

        if media_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
            doc = Document(io.BytesIO(data))
            parts = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    parts.append(" | ".join(cell.text for cell in row.cells))
            text = "\n".join(parts).strip()
            base_meta["paragraph_count"] = len(doc.paragraphs)
            base_meta["table_count"] = len(doc.tables)
            return ExtractionResult(media_type, ExtractionStatus.EXTRACTED, text or None, base_meta, "urn:anu:analyzer:docx:python-docx")

        if media_type.startswith("image/"):
            with Image.open(io.BytesIO(data)) as image:
                base_meta.update({"width": image.width, "height": image.height, "mode": image.mode, "format": image.format})
            return ExtractionResult(media_type, ExtractionStatus.METADATA_ONLY, None, base_meta, "urn:anu:analyzer:image:pillow")

        if media_type in {"audio/wav", "audio/x-wav", "audio/wave"} or filename.lower().endswith(".wav"):
            with wave.open(io.BytesIO(data), "rb") as wav:
                frames = wav.getnframes()
                rate = wav.getframerate()
                base_meta.update({
                    "channels": wav.getnchannels(),
                    "sample_rate": rate,
                    "sample_width": wav.getsampwidth(),
                    "duration_seconds": (frames / rate) if rate else None,
                })
            return ExtractionResult(media_type, ExtractionStatus.METADATA_ONLY, None, base_meta, "urn:anu:analyzer:audio:wave")

        if media_type.startswith("audio/") or media_type.startswith("video/"):
            path = stored_path
            cleanup = False
            if path is None:
                handle = tempfile.NamedTemporaryFile(delete=False, suffix=Path(filename).suffix)
                handle.write(data)
                handle.close()
                path = Path(handle.name)
                cleanup = True
            probe = _ffprobe(path)
            if cleanup:
                path.unlink(missing_ok=True)
            if probe:
                base_meta["ffprobe"] = probe
            analyzer = "urn:anu:analyzer:media:ffprobe" if probe else "urn:anu:analyzer:media:metadata-only"
            return ExtractionResult(media_type, ExtractionStatus.METADATA_ONLY, None, base_meta, analyzer)

        if media_type.startswith("text/") or filename.lower().endswith((".txt", ".md", ".csv", ".json")):
            text = data.decode("utf-8", errors="replace")
            return ExtractionResult(media_type, ExtractionStatus.EXTRACTED, text, base_meta, "urn:anu:analyzer:text:utf8")

        return ExtractionResult(media_type, ExtractionStatus.UNSUPPORTED, None, base_meta, "urn:anu:analyzer:generic")
    except Exception as exc:  # extraction failure must not destroy the immutable artifact
        base_meta["extraction_error"] = f"{type(exc).__name__}: {exc}"
        return ExtractionResult(media_type, ExtractionStatus.FAILED, None, base_meta, "urn:anu:analyzer:failed")
